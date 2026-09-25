"""Writes the exact demo commands into the speaker notes, in place.

Deliberately edits the existing file instead of regenerating the deck: slide 2
is the speaker's own and must survive untouched, and notes are the only thing
this script is allowed to change.
"""

import re
import sys
from pptx import Presentation

DECK = sys.argv[1] if len(sys.argv) > 1 else r"E:\source\repos\TDCSP2026\docs\apresentacao\Agent-Squad-TDC-SP-2026.pptx"

# Demo number -> notes. Keyed by the "DEMO n" chip on the slide rather than by
# slide index, so adding a slide earlier in the deck cannot silently move these.
NOTES = {
    1: """COMANDO EXATO — terminal na raiz do repositório da fábrica (E:\\source\\repos\\TDCSP2026):

dotnet run --project src/AgentSquad.Cli -- doctor --probe-models

Leva ~40 s, quase tudo esperando o gpt-5.5 — que é o ponto: é chamada real a cada modelo, não checagem de configuração.

ANTES DE SUBIR AO PALCO: rode uma vez "dotnet build AgentSquad.slnx -c Release" e depois acrescente "-c Release --no-build" a TODOS os comandos das demos. Sem isso você espera compilação na frente da plateia.

Se quiser provar também o Copilot (24 s, isolado de propósito):
dotnet run --project src/AgentSquad.Cli -- doctor --probe-copilot

O que apontar: cada falha traz o comando que corrige, não "algo deu errado". Repare nos 37 s do gpt-5.5 — modelo de raciocínio tem custo de latência, e isso volta no achado 2.""",

    2: """COMANDO EXATO:

dotnet run --project src/AgentSquad.Cli -- run `
  --transcript samples/meeting-transcripts/kickoff-catalogo.md `
  --repo squad-demo-catalogo --plan-only

--plan-only para depois do plano e NÃO escreve nada no GitHub. Pode repetir quantas vezes quiser, inclusive se algo der errado no palco. É a demo mais segura da apresentação.

Enquanto o agente lê a transcrição (~60 s), conte como o arquivo foi escrito: um requisito que ninguém quantificou, uma divergência que ficou sem resolução, um "obviamente" que não é óbvio, um escopo que cresceu no meio da conversa.

O teste de fogo: o agente lista o p95 como PENDÊNCIA ou registra como requisito? O gabarito está no fim do próprio arquivo da transcrição.""",

    3: """COMANDO EXATO — esta é a execução de verdade: escreve issues, branches e pull requests no GitHub.

dotnet run --project src/AgentSquad.Cli -- run `
  --request "Preciso de uma API REST de catálogo de produtos em .NET: busca paginada por nome, filtro por categoria e faixa de preço, com testes de integração." `
  --repo squad-demo-catalogo --parallel 3

SEM --plan-only (você quer que escreva) e SEM --unattended (você quer mostrar a aprovação humana do plano).

Com a interface web e o dashboard do Aspire, em outro terminal:
dotnet run --project src/AgentSquad.AppHost

Tempos: intake ~40 s · requisitos ~90 s · planejamento 4 a 6 min · agentes 3 a 20 min por onda. Você NÃO vai esperar em silêncio — o roteiro tem o que falar em cada espera.

Se o gate reprovar, COMEMORE e pare para mostrar o diagnóstico voltando para o agente.""",

    4: """COMANDO EXATO:

gh pr list --repo magoolation/squad-demo-catalogo

Para abrir um no navegador direto do terminal:
gh pr view <numero> --repo magoolation/squad-demo-catalogo --web

Ordem: abra primeiro um PR de ITEM (tabela do gate, apontamento citando TST-004 com arquivo:linha, bloco de procedência). Depois o PR de ENTREGA, o que vai de agent/run-... para main: o que entrou, o que NÃO entrou, e a última linha, "O merge é seu."

Explique os dois níveis ANTES que alguém pergunte: os PRs por item apontam para o branch de integração e servem para ler uma mudança de cada vez; o de entrega é a única coisa que propõe mexer no main. Alguns PRs de item aparecem como MERGED — o GitHub faz isso sozinho quando os commits viram alcançáveis pela base. Emende no achado 8.

Em uma frase: "revisar em pedaços, decidir de uma vez".""",
}


def demo_number(slide):
    """Returns the demo number from the slide's 'DEMO n' chip, or None."""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        match = re.match(r"^\s*DEMO\s+(\d+)\b", shape.text_frame.text)
        if match:
            return int(match.group(1))
    return None


def main():
    prs = Presentation(DECK)
    touched = {}

    for index, slide in enumerate(prs.slides, 1):
        demo = demo_number(slide)
        if demo in NOTES:
            slide.notes_slide.notes_text_frame.text = NOTES[demo]
            touched[demo] = index

    missing = sorted(set(NOTES) - set(touched))
    if missing:
        raise SystemExit(f"demos não encontradas no deck: {missing}")

    prs.save(DECK)
    for demo in sorted(touched):
        print(f"ok: DEMO {demo} -> slide {touched[demo]}")


if __name__ == "__main__":
    main()
